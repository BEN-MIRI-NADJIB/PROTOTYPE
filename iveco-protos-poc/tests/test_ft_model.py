import datetime as dt
import io
from openpyxl import Workbook
from src.excel_loader import load_excel_files
from src.ft_model import (SIT_BOTH, SIT_FAB, SIT_NONE, SIT_OK, SIT_STUDY, build_ft_table, data_warnings, ft_key, planning_table, proto_summary)

TODAY = dt.date(2026, 10, 1)
HEAD = ["Type", "FT", "Description_English", "DPT_NG", "Proto_Impact", "Check_DMU_Status", "FT_Status", "DAP_Volume_Forecast", "DAP_Volume_Sent",
        "DAP_Submission_Forecast", "Delta_Forecast_vs_Need", "Need_for_the_parts", "Material_receival_date_forecast", "New_Build_Step_PE", "Demo_Plan_B_Status", "Impact_of_delay_if_late"]
D = dt.datetime


def _wb(name, ft_rows, extra=None, sheet="FT_Impact"):
    wb = Workbook()
    ws = wb.active; ws.title = sheet
    ws.append(["Suivi prototype"])
    ws.append(HEAD)
    for r in ft_rows: ws.append(r)
    kpi = wb.create_sheet("KPI"); kpi.append(["a", "b"]); kpi.append([1, 2])  # feuille non métier
    for title, rows in (extra or {}).items():
        sh = wb.create_sheet(title)
        for r in rows: sh.append(r)
    buf = io.BytesIO(); wb.save(buf)
    return {"name": name, "content": buf.getvalue()}


ROWS = [
    ["MAIN FUNCTION", "MF", "Vehicle", None, None, None, None, None, None, None, None, None, None, None, None, None],            # hiérarchie : exclue
    ["FT", "061321", "Dashboard bracket", "BODY", "X", "KO", "Closed", 3, 3, "26W28", "LATE", D(2026, 10, 12), D(2027, 2, 1), "4. Dashboard", None, "Prevent from Starting"],
    ["FT", "061322", "Roof seal", "BODY", "X", "OK", "Closed", 1, 1, "26W37", "LATE", D(2026, 11, 9), D(2026, 12, 1), "4. Roofs panels", "Plan B validated", None],
    ["FT", "061323", "Wiring harness", "ELEC", "X", "BORDERLINE", "Open", 2, 1, "26W45", "Closed", D(2026, 11, 9), D(2026, 11, 2), "4. LV & HV harnesses", None, None],
    ["FT", "061324", "Door trim", "BODY", "X", "OK", "Closed", 1, 1, "26W37", "Closed", D(2026, 11, 9), D(2026, 11, 2), "4. Driver Area Panels", None, None],
    ["FT", "061325", "Late DAP part", "BODY", "X", "OK", "Open", 2, 1, "26W20", "IN LINE", D(2026, 11, 9), D(2026, 11, 2), "4. Seats", None, None],
    ["FT", "061399", "Hors périmètre", "BODY", "NA to P1", "KO", "Open", 1, 0, "26W20", "LATE", None, None, None, None, None],     # hors périmètre
    ["FT", "061324", "Door trim (doublon)", "BODY", "X", "OK", "Closed", 1, 1, "26W37", "Closed", D(2026, 11, 9), D(2026, 11, 2), "4. Driver Area Panels", None, None],
]


def _ft():
    extra = {"Supplier_list_LATE_with_name": [["Departement", "FT", "Supplier_name", "Action_plan"], ["BODY", 61321, "ACME 3D", "thermo + metal"]]}
    loaded, errors = load_excel_files([_wb("Follow-up_P1.xlsm", ROWS, extra)])
    assert not errors, errors
    return loaded, build_ft_table(loaded, TODAY).set_index("ft")


def test_only_business_sheets_loaded():
    loaded, _ = _ft()
    assert {s["sheet_name"] for s in loaded} == {"FT_Impact", "Supplier_list_LATE_with_name"}


def test_scope_hierarchy_and_duplicates():
    _, ft = _ft()
    assert set(ft.index) == {"61321", "61322", "61323", "61324", "61325"}  # ni MF, ni hors périmètre
    assert ft.loc["61324", "nb_lignes_source"] == 2


def test_etude_and_fabrication_split_and_situation():
    _, ft = _ft()
    assert ft.loc["61321", "etude_statut"] == "KO" and ft.loc["61321", "fab_statut"] == "EN RETARD" and ft.loc["61321", "situation"] == SIT_BOTH
    assert ft.loc["61322", "etude_statut"] == "TERMINE" and ft.loc["61322", "situation"] == SIT_FAB
    assert ft.loc["61323", "etude_statut"] == "A RISQUE" and ft.loc["61323", "situation"] == SIT_STUDY
    assert ft.loc["61324", "situation"] == SIT_OK and ft.loc["61324", "score_risque"] == 0


def test_dap_late_and_gap_and_severity():
    _, ft = _ft()
    assert ft.loc["61325", "etude_statut"] == "EN RETARD" and bool(ft.loc["61325", "dap_en_retard"])
    assert ft.loc["61321", "fab_ecart_jours"] == 112 and ft.loc["61321", "gravite_consequence"] == "BLOQUANT"
    assert ft.loc["61321", "score_risque"] > ft.loc["61322", "score_risque"] > 0
    assert "Prevent from Starting" in ft.loc["61321", "impact_consequence"]


def test_join_with_leading_zeros_and_summary():
    loaded, ft = _ft()
    assert ft.loc["61321", "fab_fournisseur"] == "ACME 3D" and ft_key("061321") == ft_key(61321) == ft_key("61321.0".split(".")[0])
    s = proto_summary(ft.reset_index()).iloc[0]
    assert s["ft_suivies"] == 5 and s["etude_expose_fab"] == 3


def test_warnings_for_prototype_without_status_and_identical_files():
    bare = [["FT", "061321", "x", "BODY", None, None, None, None, None, None, None, None, None, None, None, None]]
    sheets = []
    for n in ("Follow-up_P4.xlsm", "Follow-up_P5.xlsm"):
        # feuille FT sans colonne de statut ni proto_impact
        wb = Workbook(); ws = wb.active; ws.title = "FT_Impact"
        ws.append(["Type", "FT", "Description_English", "DPT_NG"]); ws.append(["FT", "061321", "x", "BODY"]); ws.append(["FT", "061322", "y", "ELEC"])
        buf = io.BytesIO(); wb.save(buf); sheets.append({"name": n, "content": buf.getvalue()})
    loaded, _ = load_excel_files(sheets)
    ft = build_ft_table(loaded, TODAY)
    assert set(ft["situation"]) == {SIT_NONE}
    text = " ".join(data_warnings(loaded, ft))
    assert "identiques" in text and "aucune colonne de statut" in text


def test_planning_master():
    rows = [["Departement", "Prototype", "Category", "Year", "Week", "Unit"], ["BODY", "MK32", "PLAN", 2026, 3, 2], ["BODY", "MK32", "PLAN", 2026, 4, 1], ["BODY", "MK32", "SENT", 2026, 4, 1]]
    wb = Workbook(); ws = wb.active; ws.title = "R01_Concac_Proto"
    for r in rows: ws.append(r)
    buf = io.BytesIO(); wb.save(buf)
    loaded, _ = load_excel_files([{"name": "MASTER_SUIVI_PROTO.xlsm", "content": buf.getvalue()}])
    plan = planning_table(loaded)
    assert plan[plan["category"].eq("PLAN")]["cumul"].tolist() == [2, 3]
